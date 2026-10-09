#!/usr/bin/env python3
"""Portable office inspection, isolated rendering and an editable demonstration.

No dependency installation, host-specific paths, account calls or embedded code
execution. Run with the configured interpreter; third-party imports occur only
in the explicitly requested demo command.
"""
import argparse
import hashlib
import json
import os
from pathlib import Path, PurePosixPath
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile

MAX_INPUT = 128 * 1024 * 1024
MAX_MEMBER = 32 * 1024 * 1024
MAX_EXPANDED = 256 * 1024 * 1024
MAX_FILES = 10000
MAX_PAGES = 80
TIMEOUT = 90
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      's': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
FORMATS = {'.docx', '.pptx', '.xlsx', '.pdf'}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def regular_input(value):
    path = Path(value).expanduser().absolute()
    if path.is_symlink() or not path.is_file():
        raise ValueError('Input must be an existing regular file')
    if path.suffix.lower() not in FORMATS:
        raise ValueError('Supported inputs: DOCX, PPTX, XLSX and PDF')
    if path.stat().st_size > MAX_INPUT:
        raise ValueError('Input exceeds 128 MiB limit')
    return path


def new_output(value):
    path = Path(value).expanduser().absolute()
    if '..' in path.parts:
        raise ValueError('Output may not contain parent traversal')
    # Resolve the existing parent for normal OS aliases such as macOS /var.
    parent = path.parent.resolve(strict=True)
    path = parent / path.name
    if path.exists() or path.is_symlink():
        raise ValueError('Output already exists; choose a new directory')
    path.mkdir(mode=0o700)
    return path


def executable(name):
    variable = 'HARNESS_' + name.upper()
    configured = os.environ.get(variable)
    path = shutil.which(configured or name)
    if not path:
        raise ValueError('Missing {} executable; configure {} or PATH'.format(name, variable))
    return str(Path(path).resolve())


def stop_process_tree(process):
    """Stop tool descendants as well as the wrapper on timeout/cancellation."""
    if os.name == 'nt':
        try:
            subprocess.run(['taskkill', '/PID', str(process.pid), '/T', '/F'],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                           timeout=10, check=False)
        except (OSError, subprocess.TimeoutExpired):
            pass
        if process.poll() is None:
            process.kill()
    else:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
    process.wait(timeout=10)


def run(argv, timeout=TIMEOUT):
    # File-backed logs avoid capture_output buffering unbounded renderer output.
    with tempfile.TemporaryFile() as stdout, tempfile.TemporaryFile() as stderr:
        try:
            process = subprocess.Popen(argv, stdout=stdout, stderr=stderr,
                                       start_new_session=os.name != 'nt')
            try:
                process.wait(timeout=timeout)
            except (subprocess.TimeoutExpired, KeyboardInterrupt) as error:
                stop_process_tree(process)
                if isinstance(error, KeyboardInterrupt):
                    raise
                raise ValueError('Tool timed out: ' + Path(argv[0]).name) from None
        except OSError as error:
            raise ValueError('Cannot run {}: {}'.format(Path(argv[0]).name, error)) from None
        stdout.seek(0)
        stderr.seek(0)
        out = stdout.read(65536).decode('utf-8', 'replace')
        err = stderr.read(65536).decode('utf-8', 'replace')
    if process.returncode:
        raise ValueError('Tool failed ({}): {}'.format(process.returncode, Path(argv[0]).name))
    return {'tool': Path(argv[0]).name, 'exit_code': process.returncode,
            'stdout': out, 'stderr': err}


def xml(raw):
    declaration_bytes = raw.replace(b'\x00', b'').upper()
    if b'<!DOCTYPE' in declaration_bytes or b'<!ENTITY' in declaration_bytes:
        raise ValueError('DTD/entity declarations are unsupported')
    try:
        return ET.fromstring(raw)
    except ET.ParseError:
        raise ValueError('Malformed XML package part') from None


def package(path):
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        raise ValueError('Invalid Office ZIP package') from None
    with archive:
        entries = archive.infolist()
        if len(entries) > MAX_FILES:
            raise ValueError('Office package exceeds entry limit')
        names = set()
        total = 0
        parts = {}
        for entry in entries:
            name = entry.filename
            p = PurePosixPath(name)
            if name in names or p.is_absolute() or '..' in p.parts or '\\' in name:
                raise ValueError('Duplicate or unsafe package path')
            names.add(name)
            if entry.flag_bits & 1:
                raise ValueError('Encrypted ZIP entries are unsupported')
            total += entry.file_size
            if entry.file_size > MAX_MEMBER or total > MAX_EXPANDED:
                raise ValueError('Office package exceeds expanded size limits')
            if name.endswith(('.xml', '.rels')):
                parts[name] = xml(archive.read(entry))
        if '[Content_Types].xml' not in names:
            raise ValueError('Missing Office content types')
        return names, parts


def issue(code, severity, **detail):
    return dict(code=code, severity=severity, **detail)


def inspect_office(path):
    names, parts = package(path)
    issues = []
    for declaration in parts['[Content_Types].xml']:
        if any(token in declaration.get('ContentType', '').lower()
               for token in ('macroenabled', 'vbaproject', 'activex', 'oleobject')):
            issues.append(issue('active_content_type', 'error'))
    for name in sorted(names):
        lower = name.lower()
        if any(token in lower for token in ('vbaproject', '/activex/', '/embeddings/', '/externallinks/')) or lower.endswith('/connections.xml'):
            issues.append(issue('active_or_external_package_part', 'error', part=name))
        if name.endswith('.rels'):
            for rel in parts[name]:
                if rel.get('TargetMode', '').lower() == 'external':
                    hyperlink = rel.get('Type', '').endswith('/hyperlink')
                    issues.append(issue('external_hyperlink' if hyperlink else 'external_resource',
                                        'warning' if hyperlink else 'error', part=name))
    summary = {'package_parts': len(names),
               'inspection_limit': 'ZIP/XML parsing and selected checks; complete OOXML schema and relationship integrity are not validated'}
    suffix = path.suffix.lower()
    required = {'.docx': 'word/document.xml', '.pptx': 'ppt/presentation.xml', '.xlsx': 'xl/workbook.xml'}[suffix]
    if required not in parts:
        raise ValueError('Missing primary Office XML part: ' + required)
    if suffix == '.docx':
        root = parts[required]
        summary.update(paragraphs=len(root.findall('.//w:p', NS)), tables=len(root.findall('.//w:tbl', NS)),
                       text_characters=sum(len(e.text or '') for e in root.findall('.//w:t', NS)),
                       tracked_insertions=len(root.findall('.//w:ins', NS)),
                       tracked_deletions=len(root.findall('.//w:del', NS)))
    elif suffix == '.pptx':
        slides = sorted((n for n in names if re.fullmatch(r'ppt/slides/slide[0-9]+\.xml', n)),
                        key=lambda n: int(re.search(r'slide([0-9]+)', n).group(1)))
        size = parts[required].find('p:sldSz', NS)
        if size is None:
            raise ValueError('Missing slide size')
        width, height = int(size.get('cx')), int(size.get('cy'))
        summary.update(slides=len(slides), slide_width_emu=width, slide_height_emu=height,
                       geometry_limit='Only direct, unrotated shape/picture/table bounds; no text-overflow or grouped/rotated-shape guarantee')
        for name in slides:
            tree = parts[name].find('p:cSld/p:spTree', NS)
            for shape in ([] if tree is None else list(tree)):
                transform = shape.find('p:spPr/a:xfrm', NS)
                if transform is None:
                    transform = shape.find('p:xfrm', NS)
                if transform is None or transform.get('rot', '0') != '0':
                    continue
                off, extent = transform.find('a:off', NS), transform.find('a:ext', NS)
                if off is None or extent is None:
                    continue
                x, y, cx, cy = [int(e.get(k)) for e, k in [(off, 'x'), (off, 'y'), (extent, 'cx'), (extent, 'cy')]]
                if min(x, y, cx, cy) < 0 or x + cx > width or y + cy > height:
                    issues.append(issue('off_slide_bounds', 'error', part=name, bounds_emu=[x, y, cx, cy]))
    else:
        sheets = sorted(n for n in names if re.fullmatch(r'xl/worksheets/sheet[0-9]+\.xml', n))
        formulas = 0
        for name in sheets:
            for cell in parts[name].findall('.//s:c', NS):
                formula, cache = cell.find('s:f', NS), cell.find('s:v', NS)
                if formula is not None:
                    formulas += 1
                    text = formula.text or ''
                    if re.search(r'(?:WEBSERVICE|RTD|DDE)\s*\(|\[(?:[0-9]+|[^]]+\.(?:xlsx?|xlsm|xlsb|ods))\]', text, re.I):
                        issues.append(issue('external_formula', 'error', part=name, cell=cell.get('r')))
                    if cache is None or cache.text is None:
                        issues.append(issue('formula_cache_missing_or_empty', 'warning', part=name, cell=cell.get('r')))
                if cell.get('t') == 'e':
                    issues.append(issue('spreadsheet_error', 'error', part=name, cell=cell.get('r'), value=cache.text if cache is not None else None))
        summary.update(sheets=len(sheets), formulas=formulas,
                       cache_limit='Missing/empty cache is not proof of an invalid formula; cached values can also be stale')
    return summary, issues


def inspect_data(path):
    if path.suffix.lower() == '.pdf':
        raw = path.read_bytes()
        if not raw.startswith(b'%PDF-'):
            raise ValueError('Invalid PDF header')
        issues = []
        for marker in (b'/JavaScript', b'/JS', b'/Launch', b'/EmbeddedFile'):
            if marker in raw:
                issues.append(issue('possible_active_pdf_content', 'warning'))
                break
        summary = {'bytes': len(raw), 'inspection_limit': 'Header/raw marker scan only; compressed active content and PDF structure are not validated'}
    else:
        summary, issues = inspect_office(path)
    return {'input': {'name': path.name, 'sha256': digest(path), 'format': path.suffix.lower()[1:]},
            'machine_checks': {'status': 'attention' if issues else 'passed', 'summary': summary, 'issues': issues},
            'visual_review': {'status': 'not_performed', 'required': True}}


def output_hashes(directory):
    return [{'path': p.relative_to(directory).as_posix(), 'sha256': digest(p), 'bytes': p.stat().st_size}
            for p in sorted(directory.rglob('*')) if p.is_file() and p.name != 'report.json']


def save_report(directory, report):
    report['outputs'] = output_hashes(directory)
    (directory / 'report.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8')
    return report


def pdf_pages(pdf, pdfinfo):
    result = run([pdfinfo, str(pdf)], timeout=30)
    match = re.search(r'^Pages:\s*(\d+)\s*$', result['stdout'], re.M)
    if not match:
        raise ValueError('Cannot establish PDF page count')
    pages = int(match.group(1))
    if not 1 <= pages <= MAX_PAGES:
        raise ValueError('PDF page count outside supported 1-{} range'.format(MAX_PAGES))
    return pages


def render_into(path, output):
    report = inspect_data(path)
    if any(i['severity'] == 'error' for i in report['machine_checks']['issues']):
        raise ValueError('Rendering blocked by package inspection; inspect the file first')
    soffice = executable('soffice') if path.suffix.lower() != '.pdf' else None
    pdftoppm, pdfinfo = executable('pdftoppm'), executable('pdfinfo')
    logs = []
    report['renderer'] = {'soffice': soffice, 'pdftoppm': pdftoppm, 'pdfinfo': pdfinfo}
    # No OOXML extraction and no modification of the supplied artifact.
    with tempfile.TemporaryDirectory(prefix='harness-office-') as tmp:
        temporary = Path(tmp).resolve()
        copied = temporary / ('input' + path.suffix.lower())
        shutil.copyfile(path, copied)
        profile = (temporary / 'profile').as_uri()
        if copied.suffix == '.pdf':
            pdf = output / 'preview.pdf'
            shutil.copyfile(copied, pdf)
        else:
            if copied.suffix == '.xlsx':
                recalculated = output / 'recalculated'; recalculated.mkdir()
                logs.append(run([soffice, '-env:UserInstallation=' + profile, '--headless', '--convert-to', 'xlsx', '--outdir', str(recalculated), str(copied)]))
                computed = recalculated / 'input.xlsx'
                if not computed.is_file() or not computed.stat().st_size:
                    raise ValueError('Recalculation produced no workbook')
                report['recalculated_checks'] = inspect_data(computed)['machine_checks']
                if any(i['severity'] == 'error' for i in report['recalculated_checks']['issues']):
                    raise ValueError('Recalculated workbook contains errors or external content')
                copied = computed
            logs.append(run([soffice, '-env:UserInstallation=' + profile, '--headless', '--convert-to', 'pdf', '--outdir', str(output), str(copied)]))
            pdf = output / 'input.pdf'
            if not pdf.is_file() or not pdf.stat().st_size:
                raise ValueError('Office conversion produced no PDF')
        count = pdf_pages(pdf, pdfinfo)
        logs.append(run([pdftoppm, '-f', '1', '-l', str(count), '-scale-to', '1800', '-png', str(pdf), str(output / 'page')]))
        images = sorted(output.glob('page-*.png'))
        if len(images) != count or any(not p.stat().st_size for p in images):
            raise ValueError('Rasterizer did not produce every page')
        report['render'] = {'pages': count, 'page_images': [p.name for p in images], 'max_dimension_px': 1800}
    report['tool_logs'] = logs
    report['visual_review']['pages_to_review'] = report['render']['page_images']
    if digest(path) != report['input']['sha256']:
        raise ValueError('Input changed during rendering')
    return save_report(output, report)


def inspect_command(input_path, output):
    path = regular_input(input_path)
    report = inspect_data(path)
    return save_report(new_output(output), report)


def render_command(input_path, output):
    path = regular_input(input_path)
    # Check dependencies and structure before creating an output directory.
    inspect_data(path)
    for name in (['soffice'] if path.suffix.lower() != '.pdf' else []) + ['pdftoppm', 'pdfinfo']:
        executable(name)
    return render_into(path, new_output(output))


def demo(output):
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches as WordInches, Pt as WordPt, RGBColor as WordRGB
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.oxml.xmlchemy import OxmlElement as PptxElement
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError as error:
        raise ValueError('Demo requires python-docx, python-pptx and openpyxl in the selected interpreter; missing ' + str(error.name)) from None
    for name in ['soffice', 'pdftoppm', 'pdfinfo']:
        executable(name)
    out = new_output(output)
    ink, green, paper, muted = '23372F', '28634B', 'F7F4EC', '69736D'
    rows = [('Discovery', 20, 100), ('Design', 30, 120), ('Build', 40, 150)]
    subtotal = sum(hours * rate for _, hours, rate in rows)
    expected = {'subtotal': subtotal, 'contingency': round(subtotal * .10, 2), 'total': round(subtotal * 1.10, 2)}
    d = Document(); section = d.sections[0]
    section.top_margin = WordInches(.8); section.bottom_margin = WordInches(.8)
    section.left_margin = WordInches(.85); section.right_margin = WordInches(.85)
    normal = d.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = WordPt(11)
    normal.font.color.rgb = WordRGB.from_string(ink); normal.paragraph_format.space_after = WordPt(9)
    for name in ['Title', 'Heading 1', 'Heading 2']:
        d.styles[name].font.name = 'Arial'; d.styles[name].font.color.rgb = WordRGB.from_string(green)
    for border in d.styles['Title'].element.findall('.//w:pBdr', NS):
        border.getparent().remove(border)
    section.header.paragraphs[0].text = 'FIELDWORK STUDIO  /  PROJECT BRIEF'
    section.footer.paragraphs[0].text = 'Fictional example • Planning figures in USD'
    d.add_heading('A focused website refresh', 0)
    d.add_paragraph('Decision brief • 9 October 2026')
    d.add_heading('The decision', 1)
    d.add_paragraph('Approve a four-week refresh of the studio website with a $12,760 planning budget. The work makes the service offering easier to understand and provides a clear inquiry path. All names and figures in this demonstration are fictional.')
    d.add_heading('What the project delivers', 1)
    for item in ['A concise service overview and three project stories.', 'An accessible inquiry form with clear next steps.', 'An editable content guide and a measured launch review.']:
        d.add_paragraph(item, style='List Bullet')
    d.add_heading('Budget at a glance', 1)
    table = d.add_table(rows=1, cols=3); table.style = 'Normal Table'
    for cell, value in zip(table.rows[0].cells, ['Workstream', 'Hours', 'Cost USD']): cell.text = value
    for name, hours, rate in rows:
        for cell, value in zip(table.add_row().cells, [name, str(hours), '${:,.0f}'.format(hours * rate)]): cell.text = value
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), green if row_index == 0 else paper)
            cell._tc.get_or_add_tcPr().append(shading)
            for para in cell.paragraphs:
                for run_ in para.runs:
                    run_.font.color.rgb = WordRGB.from_string('FFFFFF' if row_index == 0 else ink)
                    run_.font.bold = row_index == 0
    d.add_paragraph('Base cost $11,600 + 10% contingency $1,160 = $12,760 total.')
    d.add_page_break(); d.add_heading('Delivery and acceptance', 0)
    d.add_heading('Four weeks with a decision each week', 1)
    for title, text in [('Week 1 — Discover', 'Confirm audience, page inventory and the primary inquiry journey.'), ('Week 2 — Design', 'Review responsive layouts, navigation and the content hierarchy.'), ('Week 3 — Build', 'Implement approved pages and verify form behavior and accessibility.'), ('Week 4 — Release', 'Review the final site, publish after approval and check the inquiry flow.')]:
        d.add_heading(title, 2);d.add_paragraph(text)
    d.add_heading('Acceptance checks', 1)
    d.add_paragraph('A visitor can identify the service and send an inquiry on mobile and desktop. Keyboard navigation works, labels are readable, and the content owner can update each project story. No conversion increase is claimed before measurement.')
    d.add_heading('Boundary', 1)
    d.add_paragraph('New branding, translation and a customer portal are outside this example budget. Revisit scope if the content inventory grows beyond the agreed pages.')
    d.save(out / 'brief.docx')
    p = Presentation(); p.slide_width = Inches(13.333); p.slide_height = Inches(7.5)
    def box(slide, text, x, y, w, h, size=22, color=ink, bold=False):
        shape=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=shape.text_frame;tf.word_wrap=True
        tf.margin_left=0;tf.margin_right=0
        for n,line in enumerate(text.split('\n')):
            para=tf.paragraphs[0] if n==0 else tf.add_paragraph();para.text=line;para.font.name='Arial';para.font.size=Pt(size);para.font.bold=bold;para.font.color.rgb=RGBColor.from_string(color)
        return shape
    def slide(title, number):
        s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(paper)
        box(s,'FIELDWORK STUDIO / FICTIONAL EXAMPLE',.7,.35,11,.4,12,green,True)
        box(s,title,.7,1.05,12,1.15,34,ink,True)
        box(s,'Planning figures in USD • 9 October 2026',.7,6.95,10,.3,11,muted)
        box(s,str(number),12,6.95,.6,.3,11,muted)
        return s
    s=slide('A clearer path from interest to inquiry',1)
    box(s,'A four-week website refresh',.7,2.6,8,1,28,green)
    box(s,'Clarify the offer. Show relevant work.\nMake the next step easy to find.',.7,3.7,10,1.4,25)
    s=slide('The budget concentrates on delivery',2)
    for i,(label,amount) in enumerate([('Base cost',11600),('Contingency',1160),('Total budget',12760)]):
        x=.7+i*4.15;shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(2.7),Inches(3.8),Inches(2.4));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(green if i==2 else 'E4E9DF');shape.line.fill.background()
        shape._element.spPr.append(PptxElement('a:effectLst'))
        effect_ref = shape._element.find('p:style/a:effectRef', NS)
        if effect_ref is not None:
            effect_ref.set('idx', '0')
        color='FFFFFF' if i==2 else ink
        box(s,label,x+.25,3,3.3,.5,18,color);box(s,'${:,}'.format(amount),x+.25,3.8,3.3,.7,34,color,True)
    box(s,'90 hours across discovery, design and build; contingency is 10% of base cost.',.7,5.55,11.8,.7,18)
    s=slide('Each week ends with a concrete decision',3)
    for i,(label,desc) in enumerate([('01 Discover','Agree the journey'),('02 Design','Approve the pages'),('03 Build','Verify the behavior'),('04 Release','Review and launch')]):
        x=.7+i*3.1;box(s,label,x,2.8,2.8,.7,23,green,True);box(s,desc,x,3.7,2.8,1.2,21)
    box(s,'Accept when the inquiry flow works across devices and the content owner can update project stories.',.7,5.45,11.8,.8,20)
    p.save(out/'proposal.pptx')
    w=Workbook();ws=w.active;ws.title='Budget'
    for row in [['Fieldwork Studio project budget'],['Fictional planning example • USD'],[],['Workstream','Hours','Rate USD','Cost USD']]:ws.append(row)
    for index,(name,hours,rate) in enumerate(rows,5):ws.append([name,hours,rate,'=B{}*C{}'.format(index,index)])
    ws['C9']='Subtotal';ws['D9']='=SUM(D5:D7)';ws['C10']='Contingency';ws['B10']=.10;ws['D10']='=D9*B10';ws['C11']='Total';ws['D11']='=D9+D10'
    ws['A14']='Change hours, rates or contingency to recalculate. Figures are fictional.'
    ws.merge_cells('A1:D1');ws.merge_cells('A2:D2');ws.merge_cells('A14:D14')
    for row in ws:
        for cell in row:cell.font=Font(name='Arial',size=11,color=ink);cell.fill=PatternFill('solid',fgColor=paper);cell.alignment=Alignment(vertical='center')
    for cell in ws[4]:cell.fill=PatternFill('solid',fgColor=green);cell.font=Font(name='Arial',bold=True,color='FFFFFF')
    ws['A1'].font=Font(name='Arial',size=20,bold=True,color=green)
    for col,width in [('A',28),('B',14),('C',20),('D',20)]:ws.column_dimensions[col].width=width
    for row in range(5,12):ws.cell(row,4).number_format='"$"#,##0.00'
    for cell in ws[11]:
        cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor=green)
    ws['B10'].number_format='0%';ws.row_dimensions[1].height=34;ws.freeze_panes='B5';ws.print_area='A1:D14';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
    w.save(out/'budget.xlsx')
    artifacts=[]
    for name in ['brief.docx','proposal.pptx','budget.xlsx']:
        directory=out/(Path(name).stem+'-review');directory.mkdir()
        artifacts.append(render_into(out/name,directory))
    calculated=load_workbook(out/'budget-review/recalculated/input.xlsx',data_only=True)
    actual={'subtotal':calculated['Budget']['D9'].value,'contingency':calculated['Budget']['D10'].value,'total':calculated['Budget']['D11'].value}
    if any(abs(actual[k]-expected[k]) > .000001 for k in expected):raise ValueError('Demo recalculation does not match independent expected totals')
    if artifacts[0]['render']['pages'] != 2:raise ValueError('Demo brief did not render as two pages')
    return save_report(out,{'demo':'Fieldwork Studio fictional project','expected_totals':expected,'computed_totals':actual,
                           'machine_checks':{'status':'passed'},'artifacts':artifacts,
                           'visual_review':{'status':'not_performed','required':True}})


def interrupt_on_termination(signum, frame):
    # Let the active run() unwind its child session before CLI/temp cleanup.
    raise KeyboardInterrupt


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    commands=parser.add_subparsers(dest='command',required=True)
    for name in ['inspect','render','demo']:
        command=commands.add_parser(name)
        if name!='demo':command.add_argument('input')
        command.add_argument('--output',required=True)
    args=parser.parse_args(argv)
    previous_termination = signal.signal(signal.SIGTERM, interrupt_on_termination)
    try:
        report=demo(args.output) if args.command=='demo' else (inspect_command if args.command=='inspect' else render_command)(args.input,args.output)
        print(json.dumps(report,indent=2))
        return 1 if any(i.get('severity')=='error' for i in report['machine_checks'].get('issues',[])) else 0
    except KeyboardInterrupt:
        print('Document workflow interrupted; active tool processes stopped.', file=sys.stderr)
        return 130
    except (ValueError,OSError,zipfile.BadZipFile,KeyError,TypeError,subprocess.SubprocessError) as error:
        print('Document workflow failed: '+str(error),file=sys.stderr)
        return 2
    finally:
        signal.signal(signal.SIGTERM, previous_termination)


if __name__=='__main__':
    sys.exit(main())
